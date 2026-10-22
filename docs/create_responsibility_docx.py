import os
import sys

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    print("python-docx not found.")
    sys.exit(1)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def generate_mapped_pmo_docx():
    doc = Document()

    # Margins (0.65 in for dense, clean corporate look)
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    NAVY = RGBColor(15, 23, 42)      # #0f172a
    BLUE = RGBColor(37, 99, 235)     # #2563eb
    SLATE = RGBColor(71, 85, 105)    # #475569
    DARK = RGBColor(30, 41, 59)      # #1e293b
    WHITE = RGBColor(255, 255, 255)
    TEAL = RGBColor(13, 148, 136)    # #0d9488

    # Title & Metadata
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_org = title_p.add_run("APAR INDUSTRIES LIMITED — EPC INFRASTRUCTURE PMO\n")
    run_org.font.size = Pt(10)
    run_org.font.bold = True
    run_org.font.color.rgb = BLUE

    run_title = title_p.add_run("PMO TOOL: DOCUMENT SCOPE TO SYSTEM FEATURE MAPPING MATRIX")
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(10)
    run_sub = sub_p.add_run(
        "Complete Traceability of Management Documentation Scope to Present Software Modules & Features\n"
        "Demo Benchmarks: HTLS (GETCO 414) & OPGW (KPTCL) | SOP Ref: APAR-PMO-MAP-2026-REV4"
    )
    run_sub.font.size = Pt(9.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SLATE

    # Executive Overview
    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_after = Pt(10)
    intro_run = intro_p.add_run(
        "Executive Mapping Overview: Every document item specified in Senior Management's PMO Tool directive "
        "has been meticulously mapped to its exact frontend route, UI feature screen, interactive component, "
        "and edit/lock permissions in the present EPC Suite application. This matrix guarantees 100% feature-to-scope "
        "coverage for all 8 core management modules across both HTLS (GETCO 414) and OPGW (KPTCL) projects."
    )
    intro_run.font.size = Pt(9.5)
    intro_run.font.color.rgb = DARK

    # Section 1: Executive Mapping Table
    h1 = doc.add_heading("1. Executive Scope-to-System Feature Traceability Table", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.size = Pt(12)
        r.font.color.rgb = NAVY

    master_mapping = [
        (
            "M-01",
            "L1 / LOI / LOA / Contract Agreement / Bank Guarantees",
            "Mr. Nilesh Tiwari",
            "• L1 Intimation & Price Bid\n• LOI / Letter of Award (LOA)\n• Contract Agreement & Mother BOQ\n• Contract Performance BG (CPBG)\n• Advance Payment BG (APBG)\n• Contract Amendments & Variations",
            "• /bidding (BidDetail.tsx)\n• /award (AwardDetail.tsx)\n• /contracts (ContractList.tsx - Mother BOQ Tab)\n• /golden-period/bank-guarantees (BankGuaranteeTracker.tsx)\n• /po-amendments (POAmendmentList.tsx)",
            "Stage 0 Bidding Win/Loss, LOA Ingestion Modal, Mother BOQ Dynamic Grid, BG Validity & Claim Alerts, Scope Addendum Ledger"
        ),
        (
            "M-02",
            "Golden Activity Period (30+60 Days to Ground Break)",
            "Mr. Rakesh Sinha /\nMr. Suvesh Pahwa",
            "• 30-Day Pre-LOA Checklist\n• 60-Day Action Plan & Milestones\n• Advance Mfg Clearance (Conductor/Stubs)\n• CEIG Electrical Scheme Sanction\n• Form V & CLRA Labour License\n• Site Store Readiness Certificate\n• 60-Day Ground Break Gate Sign-off",
            "• /golden-period (GoldenPeriodDashboard.tsx)\n• /golden-period/manufacturing-clearance\n• /golden-period/statutory (StatutoryVerification.tsx)\n• /golden-period/stores (StoreSetupTracker.tsx)\n• /golden-period/insurance (InsuranceTracker.tsx)\n• /stage3 (Stage3Dashboard.tsx)",
            "60-Day Live Countdown Timer, Gate Review Checklist (BRD 2.1-2.10), Dynamic BOQ Item Mapping Cards, Statutory Document Vault, Store Infrastructure Audit, DPR-01 Lockout Gate"
        ),
        (
            "M-03",
            "Construction Dashboard Module",
            "Mr. Rakesh Sinha /\nMr. Suvesh Pahwa",
            "• WBS L2 Baseline Schedule\n• Daily Progress Reports (DPRs)\n• Quality Check Sheets (QCS)\n• Joint Measurement Certificates (JMC)\n• Pre-Commissioning Test Protocols\n• Taking Over Certificate (TOC) / Handover",
            "• /wbs (WBSManager.tsx)\n• /installation/dpr (DPRList.tsx)\n• /installation/jmc (JMCList.tsx)\n• /qa-qc/inspections (InspectionList.tsx)\n• /commissioning (CommissioningDashboard.tsx)\n• /site-surveys (SiteSurveyList.tsx)\n• /gangs (GangList.tsx)",
            "Interactive L2 Gantt Engine, Multi-gang Daily DPR Form (Civil/Erection/Stringing/OTDR), Dual-Sign JMC Attachment Modal, QCS Stage Inspection Matrix, TOC Operational Acceptance Handover"
        ),
        (
            "M-04",
            "Procurement Module",
            "Mr. Anand Raval /\nMr. Sanjay Sharma",
            "• Quotation Comparison Statement (QCS)\n• Supplier Purchase Orders (PO)\n• Factory Inspection Call Notices\n• FAT Reports & MTC Certificates\n• Client Dispatch Instructions (DI)\n• Delivery Proofs & E-Way Bills\n• Supplier & Gang Bills with Ageing",
            "• /procurement/vendors (VendorList.tsx)\n• /golden-period/vendor-finalization\n• /procurement/purchase-orders (POList.tsx)\n• /po-line-items (POLineItemList.tsx)\n• /qa-qc/inspections (InspectionList.tsx)\n• /golden-period/dispatches (DispatchTracker.tsx)\n• /inventory/dispatch (InventoryDashboard.tsx)\n• /finance/bills (BillList.tsx)",
            "QCS Vendor Evaluation Comparator, PO Creation Engine bound to Mother BOQ balance, Joint FAT Inspection Scheduler, Official DI Allocation Tracker, Transporter Proof Vault, 30/60/90-Day Bill Ageing Matrix"
        ),
        (
            "M-05",
            "Credit Control Module",
            "Mr. Arnab Bose /\nMr. Piyush Shah",
            "• Client RA Tax Invoices\n• Invoicing to Collections Mapping\n• Payment Ageing (Due/Not Due/Overdue)\n• Bill-wise & Project-wise MIS\n• Client Payment Advices & Deductions\n• Contractual Billing Milestones MIS",
            "• /collections (CollectionsDashboard.tsx)\n• /collections/invoices (InvoiceList.tsx)\n• /finance (FinanceDashboard.tsx)\n• /contracts (ContractList.tsx)",
            "Milestone Tax Invoice Generator (Supply 70%/Erection 20%/TOC 10%), Real-time Collection FIFO Mapping Ledger, Payment Ageing Brackets (Current/30/60/90+ days), Retention & LD Deduction Ledger"
        ),
        (
            "M-06",
            "Finance Module",
            "Mr. Milind Soni /\nMr. Piyush Gohil",
            "• Baseline Cost Budget (CBS)\n• Project Unbilled Revenue MIS\n• Creditors MIS & Payables Ageing\n• Bid vs Actual Profitability MIS\n• Business GFI & NFI Days Working",
            "• /finance/budgets (BudgetList.tsx)\n• /finance/bills (BillList.tsx)\n• /finance (FinanceDashboard.tsx)\n• /projects/:id (ProjectDetail.tsx)",
            "Cost Breakdown Structure (CBS) Budget Ledger, Unbilled WIP Revenue Tracker, Creditor Trade Payables Ageing Schedule, Budget vs Actual Cost Variance Analyzer, Working Capital GFI/NFI Day Computation"
        ),
        (
            "M-07",
            "Document Management System",
            "Mr. Kamal Gandhi",
            "• Time Limit Extension (EOT) Claims\n• Notice of Delay Events\n• Permit to Work (PTW Shutdowns)\n• Plant & Gang Idling Claims\n• Project Correspondence & Transmittals",
            "• /eot (EOTDashboard.tsx)\n• /eot/claims (EOTClaimList.tsx)\n• /eot/events (DelayEventList.tsx)\n• /safety (SafetyDashboard.tsx)\n• /transmittals (TransmittalList.tsx)\n• /crs (CRSList.tsx)\n• /correspondence (CorrespondenceList.tsx)",
            "Delay Impact Assessment Engine, Contractual Notice Deadline Alerter (7/14 days), Live-Line PTW Shutdown Sanction Vault, Idle Machinery Log Sheet Attacher, Incoming/Outgoing Correspondence Archive"
        ),
        (
            "M-08",
            "Inventory Management Module",
            "Mr. Gopal Nair /\nMr. Sunil Mishra",
            "• Daily Stock Ledger & Bin Cards\n• Monthly Material Reconciliation (MRR)\n• Stock Transfer Orders (STO Manifests)\n• Tools & Plants (T&P) Inventory\n• TSE Machinery Deployment & Logs\n• Scrap & Surplus Disposal Certificates",
            "• /inventory (InventoryList.tsx)\n• /inventory/dispatch (InventoryDashboard.tsx)\n• /resources (ResourceList.tsx)\n• /golden-period/stores (StoreSetupTracker.tsx)",
            "Real-time Digital Bin Cards (Receipt/Issue/Balance), 3-Way Inwarding Verification (Challan+DI+EWay), Monthly Physical vs JMC Reconciliation, Inter-project STO Authorization, Heavy TSE Machine Calibration Tracker"
        ),
    ]

    t1 = doc.add_table(rows=1, cols=6)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False
    col_w1 = [Inches(0.5), Inches(1.3), Inches(1.1), Inches(1.8), Inches(1.6), Inches(1.8)]
    headers1 = ["Mod #", "Module Name", "Lead Owner(s)", "Primary Documentation Scope", "Present System Routes & Files", "Specific Screen Features & UI Components"]

    hdr1 = t1.rows[0].cells
    for i, h in enumerate(headers1):
        hdr1[i].text = h
        hdr1[i].width = col_w1[i]
        set_cell_background(hdr1[i], "0F172A")
        set_cell_margins(hdr1[i], top=90, bottom=90, left=70, right=70)
        p = hdr1[i].paragraphs[0]
        for r in p.runs:
            r.font.size = Pt(7.5)
            r.font.bold = True
            r.font.color.rgb = WHITE

    for row_idx, row in enumerate(master_mapping):
        cells = t1.add_row().cells
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row):
            cells[i].text = val
            cells[i].width = col_w1[i]
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], top=60, bottom=60, left=70, right=70)
            p = cells[i].paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(7.5)
                if i in [0, 1]:
                    r.font.bold = True
                if i == 1:
                    r.font.color.rgb = DARK

    # Section 2: Detailed Document-by-Document Feature Mapping
    h2 = doc.add_heading("2. Granular Document-by-Document Feature & Lock Mapping", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.size = Pt(12)
        r.font.color.rgb = NAVY

    doc_details = [
        {
            "mod": "MODULE 01: COMMERCIAL & BANK GUARANTEES",
            "lead": "Mr. Nilesh Tiwari (Tendering, Contracts & Commercial)",
            "route": "/contracts, /golden-period/bank-guarantees, /award, /bidding",
            "items": [
                ("L1 Intimation & Bid Evaluation", "/bidding (BidDetail.tsx)", "Stage 0 Bid Award Card", "PDF / Scan", "Upload on L1 win; Editable until LOA issued", "Evaluates commercial ranking; initializes tender handover"),
                ("Letter of Award (LOA) / LOI", "/award (AwardDetail.tsx)", "LOA Acceptance Form", "PDF", "Upload on award receipt; LOCKED upon acceptance", "Triggers formal Project Container creation and 60-Day countdown"),
                ("Signed Contract Agreement", "/contracts (ContractList.tsx)", "Contract Master Record", "PDF", "Upload executed contract; Locked upon signing", "Formal legal basis governing all project execution obligations"),
                ("Mother BOQ Schedule", "/contracts (ContractList.tsx)", "Client BOQ Dynamic Tab", "XLSX / Grid", "Editable during contract onboarding; HARD LOCKED on sign-off", "CORE LOCK: All downstream supplier POs bound to these item balances"),
                ("Contract Performance BG (CPBG)", "/golden-period/bank-guarantees", "BG Tracker & Expiry Log", "PDF / Scan", "Upload CPBG bond; Editable for validity renewal", "Automated 60/30/15-day expiry alerts; unlocks 10% mobilization advance"),
                ("Advance Payment BG (APBG)", "/golden-period/bank-guarantees", "APBG Recovery Ledger", "PDF", "Upload APBG; Tracked for progressive recovery", "Enables client disbursement of mobilization funds"),
                ("Contract Amendments & Addendums", "/po-amendments (POAmendmentList)", "Variation Order Ledger", "PDF", "Append new revisions; original baseline preserved", "Dynamically updates BOQ items, rates, or completion deadlines"),
            ]
        },
        {
            "mod": "MODULE 02: GOLDEN ACTIVITY PERIOD (30+60 DAYS)",
            "lead": "Mr. Rakesh Sinha / Mr. Suvesh Pahwa (Project Operations & Monitoring)",
            "route": "/golden-period, /golden-period/manufacturing-clearance, /golden-period/statutory, /golden-period/stores",
            "items": [
                ("30-Day Pre-LOA Checklist", "/golden-period (Dashboard)", "Pre-LOA Risk Matrix", "XLSX", "Live editing by Rakesh Sinha / Suvesh Pahwa", "Evaluates site risks, subcontractor availability, and preliminary survey"),
                ("60-Day Master Action Plan", "/golden-period (Dashboard)", "60-Day Countdown Card", "MPP / Web", "Baseline locked; Weekly progress bar updates", "Synchronized milestone tracking to hit the ground in 60 days"),
                ("Advance Mfg Clearance", "/golden-period/manufacturing-clearance", "Clearance Allocation Card", "PDF", "Approved by Rakesh Sinha; triggers plant slots", "Allocates long-lead conductor (Apar) & buyout stubs/hardware"),
                ("CEIG Electrical Sanction", "/golden-period/statutory", "Statutory Vault (CEIG)", "PDF", "Uploaded by Statutory; Verified by Leads", "Statutory sanction for electrical layout, SLD, and corridor crossings"),
                ("Form V & CLRA Labour License", "/golden-period/statutory", "Labour Compliance Tab", "PDF", "Uploaded by HR/Statutory; Verified by Leads", "Mandatory statutory clearance to deploy contract workmen on site"),
                ("Store Readiness Certificate", "/golden-period/stores", "Store Audit Checklist", "PDF", "Uploaded by Store In-charge; Signed by Leads", "Verifies 5-acre yard, fencing, security guards, hydra crane & sleepers"),
                ("60-Day Gate Sign-off", "/golden-period (Dashboard)", "Ground Break Gate Review", "PDF / Digital", "EXECUTIVE SIGN-OFF; UNLOCKS DPR-01", "Verifies all 6 gate criteria green; authorizes physical site excavation"),
            ]
        },
        {
            "mod": "MODULE 03: CONSTRUCTION DASHBOARD",
            "lead": "Mr. Rakesh Sinha / Mr. Suvesh Pahwa (Field Operations & Quality)",
            "route": "/wbs, /installation/dpr, /installation/jmc, /qa-qc/inspections, /commissioning",
            "items": [
                ("L2 Master Baseline Schedule", "/wbs (WBSManager.tsx)", "L2 Gantt & Critical Path", "MPP / PDF", "Baseline Rev-0 locked; Revisions logged with reason", "Establishes milestone baselines and critical path tracking"),
                ("Daily Progress Reports (DPRs)", "/installation/dpr (DPRList.tsx)", "Digital Multi-Gang DPR Form", "Web / PDF", "Editable by Site Supervisor until 20:00; Locked by Lead", "Logs daily excavation, stub setting, erection, stringing, or OPGW splicing"),
                ("Quality Check Sheets (QCS)", "/qa-qc/inspections (InspectionList)", "Field QCS Inspection Matrix", "PDF / Scan", "Uploaded by QA Inspector; Non-editable once verified", "Validates pit dimensions, concrete cube strength, tower torque & sag"),
                ("Joint Measurement Certificates", "/installation/jmc (JMCList.tsx)", "JMC Dual Sign-off Modal", "PDF / Scan", "Requires dual digital signature (APAR + Utility AE)", "Verified chainage/tower measurements; prerequisite for client billing"),
                ("Pre-Commissioning Protocols", "/commissioning (Dashboard)", "Pre-Charging Protocol Card", "PDF", "Joint testing observations; Verified by Lead", "Live line pre-charging checklist, relay testing & OPGW OTDR loss"),
                ("Taking Over Certificate (TOC)", "/commissioning (Dashboard)", "TOC Commercial Handover", "PDF", "Commercial handover signed by Utility; Read-only", "Final commercial completion and transfer of asset custody to utility"),
            ]
        },
        {
            "mod": "MODULE 04: PROCUREMENT MODULE",
            "lead": "Mr. Anand Raval / Mr. Sanjay Sharma (Supply Chain & Procurement)",
            "route": "/procurement/vendors, /procurement/purchase-orders, /po-line-items, /golden-period/dispatches, /finance/bills",
            "items": [
                ("Quotation Comparison (QCS)", "/procurement/vendors (VendorList)", "QCS Evaluation Matrix", "XLSX / PDF", "Editable during evaluation; Locked upon PO approval", "Compares 3+ commercial vendor bids against tender estimates"),
                ("Supplier POs & Internal LOIs", "/procurement/purchase-orders", "PO Generator & Line Items", "PDF", "Generated in system; Locked on issue; Amendments logged", "Bounded strictly to authorized Mother BOQ item ceiling quantities"),
                ("Inspection Call Notices", "/qa-qc/inspections", "Inspection Call Scheduler", "PDF", "Uploaded by SCM; Notifies Client & QA Inspector", "Notifies readiness of manufacturing batch for formal witness testing"),
                ("FAT Reports & MTC Certificates", "/qa-qc/inspections", "FAT Mapping Card", "PDF / Scan", "Dual-signed (Vendor + Client Inspector); Uploaded by SCM", "Validates UTS, resistance, galvanizing thickness & chemical tests"),
                ("Utility Dispatch Instructions", "/golden-period/dispatches", "DI Register & Allocation", "PDF", "Uploaded by SCM/Marketing; Unlocks gate loading", "Official utility authorization required to release shipments from factory"),
                ("Delivery Proofs & E-Way Bills", "/inventory/dispatch", "Consignment Transit Vault", "PDF / Scan", "Uploaded by Logistics/Store; Matches physical receipt", "Tracks truck delivery receipts, lorry receipts (LR), and tax e-way bills"),
                ("Supplier & Gang Bills Ageing", "/finance/bills (BillList.tsx)", "Payables Ageing Grid", "PDF / Data", "Uploaded by SCM; System auto-tracks 30/60/90d ageing", "Tracks liabilities for materials & labour gangs to ensure zero site stoppage"),
            ]
        },
        {
            "mod": "MODULE 05: CREDIT CONTROL MODULE",
            "lead": "Mr. Arnab Bose / Mr. Piyush Shah (Credit Control & Commercial)",
            "route": "/collections, /collections/invoices, /finance",
            "items": [
                ("Client RA Tax Invoices", "/collections/invoices (InvoiceList)", "Milestone Tax Invoice Engine", "PDF / XML", "Generated & uploaded by Credit Control; Locked once sent", "Formatted with GST, TDS, retention deductions per contract milestone"),
                ("Invoicing-to-Collections Mapping", "/collections (Dashboard)", "Collection Reconciliation Card", "XLSX", "Live ledger mapping remittances against invoice numbers", "Reconciles gross billing against actual bank account realizations"),
                ("Payment Ageing MIS (Due/Overdue)", "/collections (Dashboard)", "Debtor Ageing Filter Tabs", "XLSX / PDF", "Auto-computed: Not Due, 1-30, 31-60, 61-90, >90 Days", "Identifies delayed utility payments for executive escalation"),
                ("Client Deductions Reconciliation", "/collections/invoices", "Deductions & LD Ledger", "PDF / Scan", "Uploaded by Credit Control; Reconciles TDS, Retention", "Tracks withheld amounts, security deposits, and liquidated damages"),
                ("Billing Milestones MIS", "/collections (Dashboard)", "Milestone Compliance Graph", "XLSX", "Tracks milestone gates (Supply 70%, Erection 20%, TOC 10%)", "Monitors realization velocity against contract billing milestones"),
            ]
        },
        {
            "mod": "MODULE 06: FINANCE MODULE",
            "lead": "Mr. Milind Soni / Mr. Piyush Gohil (Corporate Finance & Accounts)",
            "route": "/finance/budgets, /finance/bills, /finance (Dashboard)",
            "items": [
                ("Baseline Cost Budget (CBS)", "/finance/budgets (BudgetList.tsx)", "Cost Breakdown Ledger", "XLSX", "Locked baseline; Revisions require Director Approval", "Cost breakdown structure (CBS) establishing budget ceilings per work package"),
                ("Unbilled Revenue MIS (WIP)", "/finance (Dashboard)", "Unbilled Revenue Tracker", "XLSX / PDF", "Compiled monthly; verified with Site & Credit Control", "Tracks physical work completed or delivered to site pending client certification"),
                ("Creditors MIS & Payables Ageing", "/finance/bills (BillList.tsx)", "Trade Payables Ledger", "XLSX / PDF", "Accounts payable schedule; tracks vendor liabilities", "Monitors commitments and due dates for suppliers and subcontractors"),
                ("Bid vs Actual Profitability MIS", "/finance (Dashboard)", "Profitability Variance Analyzer", "XLSX / PDF", "Evaluates actual gross margin vs tender bid estimate", "Identifies cost overruns or savings across materials, labour, and logistics"),
                ("Working Capital GFI / NFI Days", "/finance (Dashboard)", "GFI & NFI Days Working Sheet", "XLSX / PDF", "Calculates Gross Funds Invested & Net Funds Invested days", "NFI Days = (Debtor + Inventory + Unbilled Days) - Creditor Days"),
            ]
        },
        {
            "mod": "MODULE 07: DOCUMENT MANAGEMENT SYSTEM",
            "lead": "Mr. Kamal Gandhi (Document Control, Contracts & Claims)",
            "route": "/eot, /eot/claims, /eot/events, /safety, /transmittals, /crs, /correspondence",
            "items": [
                ("Time Limit Extension (EOT)", "/eot/claims (EOTClaimList.tsx)", "EOT Claim Dossier Vault", "PDF / MPP", "Drafted & uploaded by Kamal Gandhi; Legal reviewed", "Formal contractual claim packages with delay impact analysis on L2 schedule"),
                ("Notice of Delay Events", "/eot/events (DelayEventList.tsx)", "Delay Event Notice Log", "PDF", "Uploaded within 48 hrs with delivery proof", "Logs site ROW hindrances, adverse weather, or client-induced stoppages"),
                ("Permit to Work (PTW Shutdowns)", "/safety (SafetyDashboard.tsx)", "PTW Grid Outage Pass", "PDF / Scan", "Uploaded by Site; Verified by Kamal Gandhi", "Utility approval for line crossing shutdowns; mandatory for stringing"),
                ("Idling Machinery & Gang Claims", "/eot/events (DelayEventList.tsx)", "Idle Plant Log Sheet", "PDF / Scan", "Uploaded from site; Verified by Kamal Gandhi", "Logs idle hours for heavy cranes, TSE machines, and gangs due to client delays"),
                ("Correspondence & Transmittals", "/transmittals, /correspondence", "Master Document Register", "PDF Archive", "Immutable logbook; auto-numbered tracking", "Central archive of all incoming/outgoing contractual letters, RFIs, and MOMs"),
            ]
        },
        {
            "mod": "MODULE 08: INVENTORY MANAGEMENT MODULE",
            "lead": "Mr. Gopal Nair / Mr. Sunil Mishra (Central Stores & Heavy Asset Logistics)",
            "route": "/inventory, /inventory/dispatch, /resources, /golden-period/stores",
            "items": [
                ("Daily Stock Ledger & Bin Cards", "/inventory (InventoryList.tsx)", "Digital Bin Card Ledger", "System / XLSX", "Real-time updates on Material Inward (MRN) & Issue (MIN)", "Tracks physical custody of conductor drums, steel stubs, and hardware"),
                ("Monthly Material Reconciliation", "/inventory (InventoryList.tsx)", "Material Balancing Matrix", "XLSX / PDF", "Compiled by Stores; Dual-verified with Site JMC", "Monthly reconciliation: Supply vs Installed vs Scrap vs Store Balance"),
                ("Stock Transfer Orders (STO)", "/inventory (InventoryList.tsx)", "Inter-Site STO Transfer Form", "PDF / Data", "Requires dual confirmation (Sending Store + Receiving Store)", "Authorizes shifting surplus materials between GETCO and KPTCL sites"),
                ("Tools & Plants (T&P) Inventory", "/resources (ResourceList.tsx)", "T&P Asset Register", "System", "Live equipment allocation and status tracking", "Tracks light tools, safety harnesses, pulling grips, and test equipment"),
                ("TSE Machinery Deployment Logs", "/resources (ResourceList.tsx)", "Heavy TSE Machinery Card", "XLSX / PDF", "Tracks Hydra Cranes, Winches, Pullers, Tensioners", "Logs machine hours, maintenance status, and calibration records"),
                ("Scrap & Surplus Disposal", "/inventory (InventoryList.tsx)", "Scrap Disposal Sign-off", "PDF / Scan", "Audited by Stores; Certified by Client Representative", "Documents return or salvage of conductor cut pieces, empty drums, and scrap"),
            ]
        }
    ]

    for m_data in doc_details:
        h3 = doc.add_heading(m_data["mod"], level=2)
        h3.paragraph_format.space_before = Pt(12)
        h3.paragraph_format.space_after = Pt(2)
        for r in h3.runs:
            r.font.size = Pt(10.5)
            r.font.color.rgb = BLUE

        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_after = Pt(4)
        r1 = meta_p.add_run("Lead Owner: ")
        r1.font.bold = True
        r1.font.size = Pt(8.5)
        r2 = meta_p.add_run(m_data["lead"] + "   |   ")
        r2.font.size = Pt(8.5)

        r3 = meta_p.add_run("System Routes: ")
        r3.font.bold = True
        r3.font.size = Pt(8.5)
        r4 = meta_p.add_run(m_data["route"])
        r4.font.size = Pt(8.5)
        r4.font.color.rgb = SLATE

        t_sub = doc.add_table(rows=1, cols=6)
        t_sub.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_sub.autofit = False
        sub_widths = [Inches(1.5), Inches(1.5), Inches(1.3), Inches(0.5), Inches(1.6), Inches(1.6)]
        sub_headers = ["Document Scope Name", "Present Route & File", "Screen Feature / Tab", "Format", "Editing & Lock Conditions", "Operational Objective & Impact"]

        hdr_sub = t_sub.rows[0].cells
        for idx, head_text in enumerate(sub_headers):
            hdr_sub[idx].text = head_text
            hdr_sub[idx].width = sub_widths[idx]
            set_cell_background(hdr_sub[idx], "1E293B")
            set_cell_margins(hdr_sub[idx], top=70, bottom=70, left=70, right=70)
            p = hdr_sub[idx].paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(7.5)
                r.font.bold = True
                r.font.color.rgb = WHITE

        for d_idx, doc_row in enumerate(m_data["items"]):
            row_cells = t_sub.add_row().cells
            bg = "F8FAFC" if d_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(doc_row):
                row_cells[c_idx].text = cell_value
                row_cells[c_idx].width = sub_widths[c_idx]
                set_cell_background(row_cells[c_idx], bg)
                set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=70, right=70)
                p = row_cells[c_idx].paragraphs[0]
                for r in p.runs:
                    r.font.size = Pt(7.5)
                    if c_idx == 0:
                        r.font.bold = True
                    elif c_idx == 1:
                        r.font.color.rgb = DARK

    # Section 3: Sign-off block
    doc.add_page_break()
    h_sign = doc.add_heading("3. Operational Governance & Lead Owner Implementation Sign-Off", level=1)
    h_sign.paragraph_format.space_before = Pt(14)
    h_sign.paragraph_format.space_after = Pt(8)
    for r in h_sign.runs:
        r.font.size = Pt(12.5)
        r.font.color.rgb = NAVY

    p_app = doc.add_paragraph()
    p_app.add_run(
        "By signing below, each designated module lead confirms that their team's required documentation scope "
        "is fully and accurately mapped to the active system routes, feature tabs, and lock rules in the In-House "
        "EPC PMO Tool for deployment on HTLS (GETCO 414) and OPGW (KPTCL) projects."
    )
    p_app.runs[0].font.size = Pt(9)

    sign_table = doc.add_table(rows=1, cols=4)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    s_widths = [Inches(1.8), Inches(2.0), Inches(1.8), Inches(1.6)]
    s_headers = ["Module Ref & Name", "Designated Lead Owner", "Signature", "Date"]

    s_hdr = sign_table.rows[0].cells
    for i, h in enumerate(s_headers):
        s_hdr[i].text = h
        s_hdr[i].width = s_widths[i]
        set_cell_background(s_hdr[i], "0F172A")
        set_cell_margins(s_hdr[i], top=80, bottom=80, left=70, right=70)
        p = s_hdr[i].paragraphs[0]
        for r in p.runs:
            r.font.size = Pt(8)
            r.font.bold = True
            r.font.color.rgb = WHITE

    sign_rows = [
        ("M-01: Commercial & BGs", "Mr. Nilesh Tiwari", "", ""),
        ("M-02: Golden Activity Period", "Mr. Rakesh Sinha / Mr. Suvesh Pahwa", "", ""),
        ("M-03: Construction Dashboard", "Mr. Rakesh Sinha / Mr. Suvesh Pahwa", "", ""),
        ("M-04: Procurement Module", "Mr. Anand Raval / Mr. Sanjay Sharma", "", ""),
        ("M-05: Credit Control Module", "Mr. Arnab Bose / Mr. Piyush Shah", "", ""),
        ("M-06: Finance Module", "Mr. Milind Soni / Mr. Piyush Gohil", "", ""),
        ("M-07: Document Management", "Mr. Kamal Gandhi", "", ""),
        ("M-08: Inventory Management", "Mr. Gopal Nair / Mr. Sunil Mishra", "", ""),
    ]

    for s_idx, s_data in enumerate(sign_rows):
        cells = sign_table.add_row().cells
        bg = "F8FAFC" if s_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(s_data):
            cells[i].text = val
            cells[i].width = s_widths[i]
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], top=120, bottom=120, left=70, right=70)
            p = cells[i].paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(8)
                if i in [0, 1]:
                    r.font.bold = True

    # Output file
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(docs_dir, "ROLE_WISE_DOCUMENT_RESPONSIBILITY_MATRIX.docx")
    doc.save(out_path)
    print(f"SUCCESS: Generated detailed feature-mapped DOCX at {out_path} (Size: {os.path.getsize(out_path)} bytes)")

if __name__ == "__main__":
    generate_mapped_pmo_docx()
